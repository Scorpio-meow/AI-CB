#!/usr/bin/env python3
import os
import sys
import argparse
from pathlib import Path
import logging
from docx import Document as DocxDocument
logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)
def convert_docx_to_txt(src_path: Path, dst_path: Path) -> None:
    try:
        doc = DocxDocument(str(src_path))
        parts = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        if not parts:
            logger.warning(f"{src_path} 無可轉換的文字內容，仍會建立空的 txt 檔案。")
        txt_content = "\n".join(parts)
        dst_path.parent.mkdir(parents=True, exist_ok=True)
        with open(dst_path, "w", encoding="utf-8") as f:
            f.write(txt_content)
        logger.info(f"轉換成功: {src_path} -> {dst_path}")
    except Exception as e:
        logger.error(f"轉換失敗: {src_path}，錯誤: {e}")
def find_docx_files(root: Path):
    for p in root.rglob("*.docx"):
        yield p
def main():
    parser = argparse.ArgumentParser(description="把 uploads 資料夾裡的 .docx 轉成 .txt")
    parser.add_argument("--uploads", "-u", default="..\\data\\uploads", help="uploads 目錄（相對於 scripts/）")
    parser.add_argument("--out", "-o", default="..\\data\\converted", help="輸出目錄（相對於 scripts/）")
    parser.add_argument("--dry-run", action="store_true", help="僅列出將要轉換的檔案，不進行轉換")
    args = parser.parse_args()
    base_dir = Path(__file__).resolve().parent
    uploads_dir = (base_dir / args.uploads).resolve()
    out_dir = (base_dir / args.out).resolve()
    if not uploads_dir.exists():
        logger.error(f"uploads 目錄不存在: {uploads_dir}")
        sys.exit(1)
    docx_files = list(find_docx_files(uploads_dir))
    if not docx_files:
        logger.info("在 uploads 資料夾中未找到任何 .docx 檔案。" )
        return
    logger.info(f"找到 {len(docx_files)} 個 .docx 檔案，輸出到 {out_dir}")
    for src in docx_files:
        relative = src.relative_to(uploads_dir)
        dst = out_dir / relative.with_suffix('.txt')
        if args.dry_run:
            logger.info(f"[DRY] {src} -> {dst}")
            continue
        convert_docx_to_txt(src, dst)
if __name__ == "__main__":
    main()
