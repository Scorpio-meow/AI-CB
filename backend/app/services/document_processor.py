#!/usr/bin/env python3
import os
import json
import csv
import logging
import hashlib
from typing import Optional
from html.parser import HTMLParser
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
import openpyxl

logger = logging.getLogger(__name__)


class _HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.result = []
        self.skip = False

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'head', 'meta', 'link'):
            self.skip = True

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'head', 'meta', 'link'):
            self.skip = False
        elif tag in ('p', 'div', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'li', 'tr', 'br', 'section', 'article'):
            self.result.append("\n")

    def handle_data(self, data):
        if not self.skip:
            text = data.strip()
            if text:
                self.result.append(text + " ")

    def get_text(self):
        return "".join(self.result).strip()


class DocumentProcessor:
    
    FILE_SIGNATURES = {
        'application/pdf': [b'%PDF'],
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document': [b'PK\x03\x04'],
        'application/vnd.openxmlformats-officedocument.presentationml.presentation': [b'PK\x03\x04'],
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': [b'PK\x03\x04'],
        'text/plain': [],
        'text/markdown': [],
        'text/x-markdown': [],
        'text/csv': [],
        'application/json': [],
        'text/json': [],
        'text/html': [],
        'text/xml': [],
        'application/xml': [],
        'text/yaml': [],
        'application/x-yaml': []
    }

    SUPPORTED_EXTENSIONS = {
        '.txt', '.md', '.markdown', '.pdf', '.docx', '.pptx', '.xlsx',
        '.csv', '.json', '.yaml', '.yml', '.xml', '.html', '.htm', '.log',
        '.py', '.js', '.ts', '.tsx', '.jsx', '.java', '.cpp', '.c', '.sql',
        '.sh', '.ini', '.env'
    }
    
    @staticmethod
    def validate_file_header(file_path: str, content_type: str) -> bool:
        signatures = DocumentProcessor.FILE_SIGNATURES.get(content_type)
        if signatures is None or not signatures:
            return True
        
        try:
            with open(file_path, 'rb') as f:
                header = f.read(512)
                
            for signature in signatures:
                if header.startswith(signature):
                    return True
            
            logger.warning(f"檔案頭部驗證失敗: {file_path}, 聲稱類型: {content_type}")
            raise ValueError(
                f"檔案頭部與聲稱的類型不符。"
                f"這可能是一個偽造的檔案或損壞的檔案。"
            )
            
        except Exception as e:
            logger.error(f"驗證檔案頭部時發生錯誤: {e}")
            raise
    
    @staticmethod
    def calculate_file_hash(file_path: str) -> str:
        sha256_hash = hashlib.sha256()
        with open(file_path, "rb") as f:
            for byte_block in iter(lambda: f.read(4096), b""):
                sha256_hash.update(byte_block)
        return sha256_hash.hexdigest()
    
    @staticmethod
    def extract_text_from_file(file_path: str, content_type: str) -> Optional[str]:
        try:
            ext = os.path.splitext(file_path)[1].lower()
            
            # 標準化 content_type
            if ext == ".docx":
                content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif ext == ".pptx":
                content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            elif ext == ".xlsx":
                content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            elif ext == ".pdf":
                content_type = "application/pdf"
            elif ext in ('.md', '.markdown'):
                content_type = "text/markdown"
            elif ext == '.csv':
                content_type = "text/csv"
            elif ext == '.json':
                content_type = "application/json"
            elif ext in ('.html', '.htm'):
                content_type = "text/html"

            DocumentProcessor.validate_file_header(file_path, content_type)
            file_size = os.path.getsize(file_path)
            max_size = int(os.getenv("MAX_FILE_SIZE_MB", "50")) * 1024 * 1024
            if file_size > max_size:
                raise ValueError(f"檔案大小 ({file_size} bytes) 超過限制")

            if ext == ".pdf" or content_type == "application/pdf":
                return DocumentProcessor._extract_from_pdf(file_path)
            elif ext == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                return DocumentProcessor._extract_from_docx(file_path)
            elif ext == ".pptx" or content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
                return DocumentProcessor._extract_from_pptx(file_path)
            elif ext == ".xlsx" or content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
                return DocumentProcessor._extract_from_xlsx(file_path)
            elif ext == ".csv" or content_type == "text/csv":
                return DocumentProcessor._extract_from_csv(file_path)
            elif ext in (".js", ".ts", ".jsx", ".tsx", ".json", ".jsonl", ".py", ".yaml", ".yml", ".ini", ".env", ".sql"):
                return DocumentProcessor._extract_from_code_or_data(file_path, ext)
            elif ext in (".html", ".htm") or content_type == "text/html":
                return DocumentProcessor._extract_from_html(file_path)
            else:
                return DocumentProcessor._extract_from_txt(file_path)
        except Exception as e:
            logger.error(f"提取文本時發生錯誤: {e}")
            raise
    
    @staticmethod
    def _extract_from_txt(file_path: str) -> str:
        encodings = ['utf-8', 'utf-8-sig', 'big5', 'gbk', 'gb2312', 'latin1']
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                if content.strip():
                    logger.info(f"成功使用 {encoding} 編碼讀取文字文件")
                    return content
            except UnicodeDecodeError:
                continue
        
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        return content

    @staticmethod
    def _extract_from_csv(file_path: str) -> str:
        encodings = ['utf-8', 'utf-8-sig', 'big5', 'gbk', 'gb2312', 'latin1']
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding, newline='') as f:
                    reader = csv.reader(f)
                    rows = []
                    for row in reader:
                        if any(cell.strip() for cell in row):
                            rows.append(" | ".join(row))
                    if rows:
                        return "\n".join(rows)
            except Exception:
                continue
        return DocumentProcessor._extract_from_txt(file_path)

    @staticmethod
    def _clean_structured_data(text: str) -> Optional[str]:
        """
        智慧識別並清洗 JS/JSON 格式的設定或資料清單 (如 const posts = [...])，
        過濾掉巨大 SVG、HTML 樣式與 Base64 等干擾噪音，重構為高密度語意文字。
        """
        import re
        # 匹配 JS 賦值陣列/物件或原生 JSON
        match = re.search(r'(?:const|let|var|module\.exports\s*=)\s*\w*\s*=?\s*(\[\s*\{[\s\S]*\}\s*\]|\{\s*\"[\s\S]*\"\s*:\s*[\s\S]*\});?', text)
        raw_json = match.group(1) if match else text.strip()
        
        # 嘗試解析 JSON
        try:
            data = json.loads(raw_json)
        except Exception:
            return None

        if isinstance(data, list):
            clean_lines = []
            noise_keys = {'embedcode', 'svg', 'icon', 'rawhtml', 'path', 'base64', 'style', 'imagedata', 'svgdata'}
            
            for idx, item in enumerate(data):
                if isinstance(item, dict):
                    parts = []
                    # 優先擷取核心欄位
                    author = item.get('author') or item.get('user') or item.get('name') or item.get('title')
                    content = item.get('content') or item.get('text') or item.get('message') or item.get('desc') or item.get('description')
                    link = item.get('postLink') or item.get('link') or item.get('url')
                    tags = item.get('tags') or item.get('categories')

                    if author:
                        parts.append(f"作者: {author}")
                    if content:
                        parts.append(f"內容: {content}")
                    if link:
                        parts.append(f"連結: {link}")
                    if tags:
                        tags_str = ', '.join(tags) if isinstance(tags, list) else str(tags)
                        parts.append(f"標籤: {tags_str}")
                    
                    # 擷取其餘非噪音欄位
                    for k, v in item.items():
                        if k.lower() in noise_keys:
                            continue
                        if k in ('author', 'user', 'name', 'title', 'content', 'text', 'message', 'desc', 'description', 'postLink', 'link', 'url', 'tags', 'categories'):
                            continue
                        if isinstance(v, str):
                            # 過濾過長 SVG/HTML 字串
                            if v.startswith('<blockquote') or v.startswith('<svg') or (len(v) > 500 and ('<path' in v or 'data:image' in v)):
                                continue
                            if v.strip():
                                parts.append(f"{k}: {v.strip()}")
                        elif v is not None and not isinstance(v, (dict, list)):
                            parts.append(f"{k}: {v}")

                    if parts:
                        clean_lines.append(f"【記錄 {idx + 1}】\n" + "\n".join(parts))
                elif isinstance(item, str) and item.strip():
                    clean_lines.append(item.strip())

            if clean_lines:
                logger.info(f"結構化資料降噪解析成功: 從陣列提取出 {len(clean_lines)} 條語意記錄")
                return "\n\n---\n\n".join(clean_lines)

        elif isinstance(data, dict):
            # 若為字典物件，展開鍵值對
            clean_lines = []
            for k, v in data.items():
                if isinstance(v, str) and len(v) < 2000 and not v.startswith('<svg'):
                    clean_lines.append(f"{k}: {v}")
                elif isinstance(v, (int, float, bool)):
                    clean_lines.append(f"{k}: {v}")
            if clean_lines:
                return "\n".join(clean_lines)

        return None

    @staticmethod
    def _extract_from_code_or_data(file_path: str, ext: str) -> str:
        raw_text = DocumentProcessor._extract_from_txt(file_path)
        
        # 1. 嘗試以結構化資料解析（針對包含大量陣列/物件的 config/data 檔案）
        cleaned = DocumentProcessor._clean_structured_data(raw_text)
        if cleaned:
            return cleaned
        
        # 2. 若為一般程式碼，過濾超長 Base64 / SVG 內嵌字串避免向量分塊爆炸
        import re
        cleaned_code = re.sub(r'data:image\/[a-zA-Z]+;base64,[a-zA-Z0-9+/=]{100,}', '[BASE64_IMAGE_OMITTED]', raw_text)
        cleaned_code = re.sub(r'<svg[\s\S]*?<\/svg>', '[SVG_ICON_OMITTED]', cleaned_code)
        return cleaned_code

    @staticmethod
    def _extract_from_html(file_path: str) -> str:
        raw_text = DocumentProcessor._extract_from_txt(file_path)
        try:
            parser = _HTMLTextExtractor()
            parser.feed(raw_text)
            extracted = parser.get_text()
            if extracted:
                return extracted
        except Exception as e:
            logger.warning(f"HTML 解析失敗，改為純文字讀取: {e}")
        return raw_text
    
    @staticmethod
    def _extract_from_pdf(file_path: str) -> str:
        from app.services.pdf_service import PDFService
        return PDFService.extract_text_robust(file_path)
    
    @staticmethod
    def _extract_from_docx(file_path: str) -> str:
        doc = DocxDocument(file_path)
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text)
        
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_content.append(" | ".join(row_text))
        
        if not text_content:
            raise ValueError("DOCX 文件中沒有可提取的文本內容")
        
        full_text = "\n".join(text_content)
        logger.info(f"成功從 DOCX 提取 {len(doc.paragraphs)} 個段落和 {len(doc.tables)} 個表格，共 {len(full_text)} 字符")
        return full_text

    @staticmethod
    def _extract_from_pptx(file_path: str) -> str:
        prs = Presentation(file_path)
        slides_text = []
        for idx, slide in enumerate(prs.slides):
            slide_lines = [f"--- 投影片 {idx + 1} ---"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_lines.append(text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            slide_lines.append(" | ".join(row_cells))
            if len(slide_lines) > 1:
                slides_text.append("\n".join(slide_lines))
        
        if not slides_text:
            raise ValueError("PPTX 文件中沒有可提取的文本內容")
        
        full_text = "\n\n".join(slides_text)
        logger.info(f"成功從 PPTX 提取 {len(prs.slides)} 頁投影片，共 {len(full_text)} 字符")
        return full_text

    @staticmethod
    def _extract_from_xlsx(file_path: str) -> str:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets_text = []
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            sheet_lines = [f"--- 工作表: {sheetname} ---"]
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() != "" for cell in row):
                    row_str = [str(c).strip() if c is not None else "" for c in row]
                    sheet_lines.append(" | ".join(row_str))
            if len(sheet_lines) > 1:
                sheets_text.append("\n".join(sheet_lines))
        
        if not sheets_text:
            raise ValueError("XLSX 文件中沒有可提取的文本內容")
        
        full_text = "\n\n".join(sheets_text)
        logger.info(f"成功從 XLSX 提取 {len(wb.sheetnames)} 個工作表，共 {len(full_text)} 字符")
        return full_text
    
    @staticmethod
    def validate_file_type(content_type: str, file_path: str = None) -> bool:
        if content_type in DocumentProcessor.FILE_SIGNATURES:
            return True
        if file_path:
            ext = os.path.splitext(file_path)[1].lower()
            if ext in DocumentProcessor.SUPPORTED_EXTENSIONS:
                return True
        return False
    
    @staticmethod
    def get_file_info(file_path: str, content_type: str) -> dict:
        try:
            stat = os.stat(file_path)
            return {
                "file_size": stat.st_size,
                "content_type": content_type,
                "is_supported": DocumentProcessor.validate_file_type(content_type, file_path)
            }
        except Exception as e:
            logger.error(f"獲取文件信息時發生錯誤: {e}")
            return {"error": str(e)}

    @classmethod
    def generate_document_summary(cls, filename: str, content: str, file_type: str = "") -> str:
        """
        純粹根據提取後之實際文本結構、章節大綱與欄位特徵，動態推導出專屬之高密度主題與功能描述（拒絕寫死任何檔名）。
        """
        if not content or not content.strip():
            return f"收錄內部文件《{filename}》。"

        lines = [line.strip() for line in content.strip().splitlines() if line.strip()]
        if not lines:
            return f"收錄內部文件《{filename}》。"

        # 1. 偵測結構化記錄庫（例如 【記錄 N】...）
        record_count = content.count("【記錄 ")
        if record_count > 0:
            first_block = content[:1500]
            fields = []
            for field_name in ["作者", "內容", "連結", "標籤", "時間", "標題", "狀態", "金額", "說明"]:
                if f"{field_name}:" in first_block or f"{field_name}：" in first_block:
                    fields.append(field_name)
            fields_str = "、".join(fields) if fields else "多項屬性欄位"
            return f"收錄約 {record_count} 筆結構化社群/資料記錄（包含欄位：{fields_str}），支援依據作者帳號、內容文字、特定網址 URL 與主題標籤進行精準檢索。"

        # 2. 偵測 Markdown 標題結構或章節大綱
        headers = []
        for line in lines[:50]:
            if line.startswith("#"):
                clean_header = line.lstrip("#").strip()
                if clean_header and len(clean_header) > 1 and clean_header not in headers:
                    headers.append(clean_header)
            elif any(line.startswith(prefix) for prefix in ("【", "第", "一、", "二、", "三、", "1.", "2.")):
                clean_line = line[:40].strip()
                if clean_line and clean_line not in headers:
                    headers.append(clean_line)

        if headers:
            main_title = headers[0]
            sub_topics = "、".join(headers[1:5])
            if sub_topics:
                return f"主題為「{main_title}」，涵蓋章節與重點包括：{sub_topics} 等專業資訊。"
            else:
                snippet = " ".join([l for l in lines[1:5] if not l.startswith("---")])[:80]
                return f"主題為「{main_title}」（重點摘要：{snippet}...）。"

        # 3. 偵測表格或 CSV 欄位特徵
        if "," in lines[0] or "\t" in lines[0] or "|" in lines[0]:
            cols = [c.strip() for c in lines[0].replace("|", ",").split(",") if c.strip()]
            if len(cols) >= 2:
                cols_str = "、".join(cols[:8])
                return f"結構化表格數據，收錄約 {len(lines) - 1} 筆資料，包含欄位：{cols_str}。"

        # 4. 通用文本摘要
        text_snippet = " ".join(lines[:4])[:120].replace("  ", " ")
        return f"收錄關於「{text_snippet}...」之內部專業文件資料。"
